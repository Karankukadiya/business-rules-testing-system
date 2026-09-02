"""Generates the Excel deliverables in output/.

Usage:
    python report_generator.py            # create any MISSING files (never overwrites tester results)
    python report_generator.py --force    # regenerate ALL files from scratch (erases recorded results!)
    python report_generator.py --summary  # rebuild test_summary.xlsx from the current suite + defect log
"""
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from validation import BUSINESS_RULES

OUTPUT_DIR = Path(__file__).parent / "output"
SUITE_FILE = OUTPUT_DIR / "manual_test_suite.xlsx"
DEFECT_FILE = OUTPUT_DIR / "defect_log.xlsx"
SUMMARY_FILE = OUTPUT_DIR / "test_summary.xlsx"

TEST_HEADERS = ["Test Case ID", "Requirement ID", "Test Scenario", "Preconditions", "Test Steps",
                "Test Data", "Expected Result", "Actual Result", "Status", "Priority", "Test Type", "Remarks"]
DEFECT_HEADERS = ["Defect ID", "Related Test Case ID", "Requirement ID", "Defect Title", "Steps to Reproduce",
                  "Expected Result", "Actual Result", "Severity", "Priority", "Status", "Reported Date",
                  "Retest Result"]
PENDING = "Pending manual execution"
FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
THIN = Side(style="thin", color="BFBFBF")

# (Requirement, Scenario, Preconditions, Steps, Test Data, Expected, Priority, Type)
DB = "App running; employee database is empty or does not contain the test ID."
VAL = "App running; Business Rule Validation page open (validation-only, nothing is saved)."
ADD = "App running; Employee Management > Add Employee form open."
RAW = [
    ("BR-01", "Add employee with a valid unique ID", ADD, "Fill all fields with valid data; click Save.",
     "ID=EMP-001, Name=Asha Patel, Age=30, Salary=50000, IT, Full-time, 2023-01-15, Active",
     "Record saved; success message; employee appears in list.", "High", "Positive"),
    ("BR-01", "Reject empty Employee ID", ADD, "Leave Employee ID empty; fill other fields validly; click Save.",
     "ID=(empty)", "Not saved; error names 'employee_id' and says an ID is required.", "High", "Negative"),
    ("BR-01", "Reject whitespace-only Employee ID", ADD, "Enter spaces only in Employee ID; click Save.",
     "ID='   '", "Not saved; error names 'employee_id' (required).", "Medium", "Negative"),
    ("BR-01", "Reject duplicate Employee ID", "Employee EMP-001 already saved.",
     "Add another employee using ID EMP-001; click Save.", "ID=EMP-001, Name=Ravi Shah, other fields valid",
     "Not saved; duplicate-ID error; only one EMP-001 remains.", "High", "Duplicate"),
    ("BR-01", "Duplicate ID detected after trimming spaces", "Employee EMP-001 already saved.",
     "Add employee with ID ' EMP-001 ' (leading/trailing spaces); click Save.", "ID=' EMP-001 '",
     "Not saved; duplicate-ID error (spaces are trimmed before comparing).", "Medium", "Boundary"),
    ("BR-02", "Reject empty employee name", ADD, "Leave Name empty; click Save.", "Name=(empty)",
     "Not saved; error names 'name' and asks for a full name.", "High", "Negative"),
    ("BR-02", "Reject whitespace-only name", ADD, "Enter only spaces in Name; click Save.", "Name='    '",
     "Not saved; error names 'name' (cannot be only spaces).", "High", "Negative"),
    ("BR-02", "Accept name with surrounding spaces (trimmed)", ADD, "Enter '  Meera Nair  ' as Name; save.",
     "Name='  Meera Nair  '", "Saved; stored name is 'Meera Nair' without extra spaces.", "Low", "Positive"),
    ("BR-03", "Age 17 rejected (just below minimum)", VAL, "Enter valid data with Age 17; click Validate.",
     "Age=17", "BR-03 shows FAIL; message says age must be 18 to 60.", "High", "Boundary"),
    ("BR-03", "Age 18 accepted (minimum)", VAL, "Enter valid data with Age 18; click Validate.", "Age=18",
     "BR-03 shows PASS.", "High", "Boundary"),
    ("BR-03", "Age 60 accepted (maximum)", VAL, "Enter valid data with Age 60; click Validate.", "Age=60",
     "BR-03 shows PASS.", "High", "Boundary"),
    ("BR-03", "Age 61 rejected (just above maximum)", VAL, "Enter valid data with Age 61; click Validate.",
     "Age=61", "BR-03 shows FAIL; message says age must be 18 to 60.", "High", "Boundary"),
    ("BR-03", "Non-numeric age rejected", VAL, "Enter 'abc' as Age; click Validate.", "Age=abc",
     "BR-03 shows FAIL; message says age must be a whole number.", "Medium", "Negative"),
    ("BR-04", "Salary 0 rejected", VAL, "Enter Salary 0; click Validate.", "Salary=0",
     "BR-04 shows FAIL; salary must be above 0.", "High", "Boundary"),
    ("BR-04", "Salary 0.01 accepted (just above zero)", VAL, "Enter Salary 0.01; click Validate.",
     "Salary=0.01", "BR-04 shows PASS.", "Medium", "Boundary"),
    ("BR-04", "Salary 200000 accepted (maximum)", VAL, "Enter Salary 200000; click Validate.",
     "Salary=200000", "BR-04 shows PASS.", "High", "Boundary"),
    ("BR-04", "Salary 200001 rejected (above maximum)", VAL, "Enter Salary 200001; click Validate.",
     "Salary=200001", "BR-04 shows FAIL; salary must not exceed 200,000.", "High", "Boundary"),
    ("BR-04", "Negative salary rejected", VAL, "Enter Salary -500; click Validate.", "Salary=-500",
     "BR-04 shows FAIL.", "Medium", "Negative"),
    ("BR-04", "Non-numeric salary rejected", VAL, "Enter 'fifty' as Salary; click Validate.",
     "Salary=fifty", "BR-04 shows FAIL; message says salary must be a number.", "Medium", "Negative"),
    ("BR-05", "Valid department accepted", VAL, "Enter department 'Operations'; click Validate.",
     "Department=Operations", "BR-05 shows PASS.", "Medium", "Positive"),
    ("BR-05", "Invalid department rejected", VAL, "Enter department 'Marketing'; click Validate.",
     "Department=Marketing", "BR-05 shows FAIL and lists the allowed departments.", "High", "Negative"),
    ("BR-05", "Empty department rejected", VAL, "Leave department empty; click Validate.",
     "Department=(empty)", "BR-05 shows FAIL.", "Medium", "Negative"),
    ("BR-06", "Valid employment type accepted", VAL, "Enter type 'Contract'; click Validate.",
     "Employment type=Contract", "BR-06 shows PASS.", "Medium", "Positive"),
    ("BR-06", "Invalid employment type rejected", VAL, "Enter type 'Intern'; click Validate.",
     "Employment type=Intern", "BR-06 shows FAIL and lists the allowed types.", "High", "Negative"),
    ("BR-07", "Joining date today accepted", VAL, "Enter today's date; click Validate.",
     "Joining date=today", "BR-07 shows PASS.", "High", "Boundary"),
    ("BR-07", "Joining date tomorrow rejected", VAL, "Enter tomorrow's date; click Validate.",
     "Joining date=tomorrow", "BR-07 shows FAIL; date is in the future.", "High", "Boundary"),
    ("BR-07", "Impossible date rejected", VAL, "Enter 2024-02-30; click Validate.",
     "Joining date=2024-02-30", "BR-07 shows FAIL; asks for YYYY-MM-DD.", "Medium", "Negative"),
    ("BR-08", "Valid status 'On Leave' accepted", VAL, "Enter status 'On Leave'; click Validate.",
     "Status=On Leave", "BR-08 shows PASS.", "Medium", "Positive"),
    ("BR-08", "Invalid status rejected", VAL, "Enter status 'Retired'; click Validate.",
     "Status=Retired", "BR-08 shows FAIL and lists allowed statuses.", "High", "Negative"),
    ("BR-09", "Invalid record is not stored in database",
     "Database empty; DB tool (DB Browser/sqlite3) available.",
     "Try to save a record with Age 61; then run SQL query Q3 and 'SELECT COUNT(*) FROM employees'.",
     "ID=EMP-900, Age=61, others valid", "Save rejected; table still has 0 rows; Q3 returns no rows.",
     "High", "Database Integrity"),
    ("BR-09", "Multiple invalid fields all reported, nothing saved", ADD,
     "Enter empty name, age 5, salary 0, then Save.", "Name=(empty), Age=5, Salary=0",
     "Not saved; errors listed for name, age and salary.", "High", "Error Handling"),
    ("BR-10", "Edit name only; other records unchanged", "EMP-001 and EMP-002 saved.",
     "Edit EMP-001 name to 'Asha P.'; save; view list.", "EMP-001 new name='Asha P.'",
     "Only EMP-001 changed; EMP-002 identical to before.", "High", "CRUD"),
    ("BR-10", "Changing ID to another employee's ID is rejected", "EMP-001 and EMP-002 saved.",
     "Edit EMP-002 and change its ID to EMP-001; save.", "New ID=EMP-001",
     "Update rejected with duplicate-ID error; both records unchanged.", "High", "Duplicate"),
    ("BR-10", "Saving edit without changing ID gives no false duplicate error", "EMP-001 saved.",
     "Edit EMP-001 salary only; save.", "Salary=60000", "Saved successfully; no duplicate error.",
     "High", "Regression"),
    ("BR-11", "Delete without confirmation does nothing", "EMP-001 saved.",
     "Select EMP-001; click Delete without ticking confirmation.", "EMP-001",
     "Record NOT deleted; message asks for confirmation.", "High", "Negative"),
    ("BR-11", "Delete with confirmation removes only that record", "EMP-001 and EMP-002 saved.",
     "Select EMP-001; tick confirmation; click Delete; check list.", "EMP-001",
     "EMP-001 removed; EMP-002 remains.", "High", "CRUD"),
    ("BR-11", "Cancelling delete keeps the record", "EMP-001 saved.",
     "Click Delete, then leave confirmation unticked and navigate away; return to list.", "EMP-001",
     "EMP-001 is still present.", "Medium", "Negative"),
    ("BR-12", "Error message names the invalid field", ADD, "Save a record with an invalid age.",
     "Age=99", "Message clearly includes the field name (age).", "Medium", "Error Handling"),
    ("BR-12", "Error message explains how to fix it", ADD, "Save a record with salary 0.",
     "Salary=0", "Message states the valid range / what to enter.", "Medium", "Error Handling"),
    ("BR-01, BR-09, BR-10, BR-11", "Regression: add, search, edit, delete cycle after a fix", ADD,
     "Add EMP-001; search by ID; edit name; delete with confirmation; run SQL Q1-Q8.",
     "EMP-001 valid data", "All steps work; Q1-Q8 return no rows.", "High", "Regression"),
]
TEST_CASES = [(f"TC-{i:03d}",) + row for i, row in enumerate(RAW, 1)]

SAMPLE_DEFECTS = [
    ("SAMPLE-DEF-001", "TC-012", "BR-03", "[SAMPLE] Age 61 accepted as valid",
     "1. Open Business Rule Validation. 2. Enter valid data with Age 61. 3. Click Validate.",
     "BR-03 FAIL; age must be 18-60.", "BR-03 shows PASS (hypothetical upper bound coded as <= 61).",
     "High", "High", "Open", "Not retested"),
    ("SAMPLE-DEF-002", "TC-033", "BR-10", "[SAMPLE] Editing an ID to an existing ID overwrites another record",
     "1. Save EMP-001 and EMP-002. 2. Edit EMP-002, change ID to EMP-001. 3. Save.",
     "Update rejected: duplicate ID.", "Update accepted; EMP-001 data replaced (hypothetical).",
     "Critical", "High", "Open", "Not retested"),
    ("SAMPLE-DEF-003", "TC-007", "BR-02", "[SAMPLE] Whitespace-only name saved",
     "1. Open Add Employee. 2. Enter '   ' as Name, other fields valid. 3. Save.",
     "Rejected with error on 'name'.", "Record saved with blank-looking name (hypothetical).",
     "Medium", "Medium", "Open", "Not retested"),
    ("SAMPLE-DEF-004", "TC-035", "BR-11", "[SAMPLE] Delete works without confirmation",
     "1. Save EMP-001. 2. Click Delete without ticking confirmation.",
     "Record kept; confirmation required.", "Record deleted immediately (hypothetical).",
     "High", "High", "Fixed - Pending Retest", "Not retested"),
]


def _style_sheet(ws, widths, header_row=1):
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for cell in ws[header_row]:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = Font(name=FONT, size=10, bold=cell.font.bold)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = ws.dimensions


def create_test_suite(path=SUITE_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    ws.append(TEST_HEADERS)
    for tc_id, req, scen, pre, steps, data, exp, prio, ttype in TEST_CASES:
        ws.append([tc_id, req, scen, pre, steps, data, exp, PENDING, "Not Executed", prio, ttype, PENDING])
    _style_sheet(ws, [11, 14, 34, 32, 44, 34, 40, 24, 14, 10, 18, 24])
    last = len(TEST_CASES) + 1
    dv = DataValidation(type="list", formula1='"Not Executed,Pass,Fail,Blocked"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{last}")
    for text, color in (("Pass", "C6EFCE"), ("Fail", "FFC7CE"), ("Blocked", "FFEB9C")):
        ws.conditional_formatting.add(f"I2:I{last}", CellIsRule(
            operator="equal", formula=[f'"{text}"'], fill=PatternFill("solid", bgColor=color)))

    tr = wb.create_sheet("Requirements Traceability")
    tr.append(["Requirement ID", "Requirement Description", "Test Case IDs", "Number of Test Cases"])
    for n, (rule, desc) in enumerate(BUSINESS_RULES.items(), 2):
        ids = [t[0] for t in TEST_CASES if rule in [r.strip() for r in t[1].split(",")]]
        tr.append([rule, desc, ", ".join(ids), f"=COUNTIF('Test Cases'!B$2:B${last},\"*\"&A{n}&\"*\")"])
    _style_sheet(tr, [14, 62, 46, 20])
    wb.save(path)


def create_defect_log(path=DEFECT_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Defects"
    ws.append(DEFECT_HEADERS)
    for d in SAMPLE_DEFECTS:
        ws.append(list(d[:10]) + [date.today(), d[10]])
    _style_sheet(ws, [17, 14, 13, 38, 46, 30, 38, 11, 10, 22, 14, 16])
    for row in ws.iter_rows(min_row=2, min_col=11, max_col=11):
        row[0].number_format = "yyyy-mm-dd"
    note = wb.create_sheet("Read Me")
    for line in ["IMPORTANT: Defects with IDs starting 'SAMPLE-' are ILLUSTRATIVE EXAMPLES.",
                 "They were NOT found by actually executing tests. Their 'actual results' are hypothetical.",
                 "Replace or delete them and log real defects (ID like DEF-001) after genuine manual execution."]:
        note.append([line])
    note.column_dimensions["A"].width = 100
    for cell in note["A"]:
        cell.font = Font(name=FONT, bold=True, color="C00000")
    wb.save(path)


def _read_column(path, sheet, column_index):
    ws = load_workbook(path, data_only=True)[sheet]
    return [r[column_index] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]


def status_counts(suite_path=SUITE_FILE):
    """Count Pass/Fail/Blocked/Not Executed in the suite as currently saved (used by the app)."""
    counts = {"Total": 0, "Pass": 0, "Fail": 0, "Blocked": 0, "Not Executed": 0}
    for status in _read_column(suite_path, "Test Cases", 8):
        counts["Total"] += 1
        if status in counts:
            counts[status] += 1
    executed = counts["Pass"] + counts["Fail"]
    counts["Pass %"] = round(100 * counts["Pass"] / executed, 1) if executed else 0.0
    return counts


def create_summary(suite_path=SUITE_FILE, defect_path=DEFECT_FILE, path=SUMMARY_FILE):
    """Copies current Status values from the suite/defect log into data sheets; summary uses formulas."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Summary"
    suite = load_workbook(suite_path, data_only=True)["Test Cases"]
    td = wb.create_sheet("Test Data")
    td.append(["Test Case ID", "Requirement ID", "Status"])
    for r in suite.iter_rows(min_row=2, values_only=True):
        if r[0]:
            td.append([r[0], r[1], r[8]])
    dd = wb.create_sheet("Defect Data")
    dd.append(["Defect ID", "Status"])
    for r in load_workbook(defect_path, data_only=True)["Defects"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            dd.append([r[0], r[9]])
    _style_sheet(td, [14, 26, 16])
    _style_sheet(dd, [18, 26])
    T = "'Test Data'!"
    rows = [
        ("Total test cases", f"=COUNTA({T}A2:A500)"),
        ("Passed", f'=COUNTIF({T}C2:C500,"Pass")'),
        ("Failed", f'=COUNTIF({T}C2:C500,"Fail")'),
        ("Blocked", f'=COUNTIF({T}C2:C500,"Blocked")'),
        ("Not Executed", f'=COUNTIF({T}C2:C500,"Not Executed")'),
        ("Pass percentage of executed tests", "=IF((B3+B4)=0,0,B3/(B3+B4))"),
        ("Open sample defects", "=COUNTIF('Defect Data'!B2:B500,\"Open\")"),
        ("Requirements covered", '=COUNTIF(B14:B25,">0")&" of 12"'),
    ]
    ws.append(["Metric", "Value"])
    for label, formula in rows:
        ws.append([label, formula])
    ws["B7"].number_format = "0.0%"
    ws["A10"] = "Notes: 'Executed' = Pass + Fail. Pass % shows 0.0% until tests are executed. Open defects are SAMPLE data."
    ws["A10"].font = Font(name=FONT, italic=True, size=9)
    ws["A12"] = "Requirement coverage"
    ws["A12"].font = Font(name=FONT, bold=True)
    ws["A13"], ws["B13"] = "Requirement ID", "Test cases"
    for n, rule in enumerate(BUSINESS_RULES, 14):
        ws[f"A{n}"] = rule
        ws[f"B{n}"] = f'=COUNTIF({T}B2:B500,"*"&A{n}&"*")'
    _style_sheet(ws, [38, 20])
    ws.auto_filter.ref = None
    for cell in ws[13]:
        cell.font = Font(name=FONT, bold=True)
    wb.save(path)


def generate_all(force=False):
    OUTPUT_DIR.mkdir(exist_ok=True)
    created = []
    if force or not SUITE_FILE.exists():
        create_test_suite(); created.append(SUITE_FILE.name)
    if force or not DEFECT_FILE.exists():
        create_defect_log(); created.append(DEFECT_FILE.name)
    create_summary(); created.append(SUMMARY_FILE.name)
    return created


if __name__ == "__main__":
    if "--summary" in sys.argv:
        create_summary(); print("Rebuilt", SUMMARY_FILE)
    else:
        print("Generated:", ", ".join(generate_all(force="--force" in sys.argv)))
