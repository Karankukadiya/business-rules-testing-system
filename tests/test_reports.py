from openpyxl import load_workbook

import ai_assistant
import report_generator as rg


def test_suite_has_40_unique_cases_covering_all_rules():
    ids = [t[0] for t in rg.TEST_CASES]
    assert len(ids) == 40 == len(set(ids))
    covered = {r.strip() for t in rg.TEST_CASES for r in t[1].split(",")}
    assert covered == {f"BR-{n:02d}" for n in range(1, 13)}


def test_generated_workbooks_have_expected_sheets(tmp_path):
    suite, defects, summary = tmp_path / "s.xlsx", tmp_path / "d.xlsx", tmp_path / "t.xlsx"
    rg.create_test_suite(suite)
    rg.create_defect_log(defects)
    rg.create_summary(suite, defects, summary)
    wb = load_workbook(suite)
    assert wb.sheetnames == ["Test Cases", "Requirements Traceability"]
    assert wb["Test Cases"].max_row == 41
    assert {r[8] for r in wb["Test Cases"].iter_rows(min_row=2, values_only=True)} == {"Not Executed"}
    assert load_workbook(defects)["Defects"].max_row - 1 >= 3
    assert "Test Summary" in load_workbook(summary).sheetnames


def test_ai_fallback_drafts_and_summarises():
    assert len(ai_assistant.draft_scenarios("BR-03")) >= 4
    assert ai_assistant.summarize_defects(["BR-03 age 61 accepted"])["total"] == 1
    assert ai_assistant.summarize_defects([])["total"] == 0
